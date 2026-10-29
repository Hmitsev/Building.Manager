import io
import os
import smtplib
from email.message import EmailMessage
from datetime import date, datetime, time
from decimal import Decimal
from urllib.parse import quote
from pathlib import Path

import pandas as pd
from PIL import Image
from openpyxl import load_workbook
import streamlit as st

from auth import hash_password, verify_password, login_user, logout_user, current_user
from config import APP_NAME, CURRENCY
from database import init_db, fetch_all, fetch_one, execute
from services import MONTHS, create_year_obligations, obligation_rows, create_payment, audit

st.set_page_config(page_title=APP_NAME, page_icon="🏢", layout="wide", initial_sidebar_state="expanded")

# PERFORMANCE SAFETY:
# - schema initialization is cached per server process
# - only the selected admin section is rendered
# Business rules, payment allocation, historical import and audit logic stay unchanged.
st.markdown("""
<style>
:root{--line:rgba(148,163,184,.16);--text:#f8fafc}.stApp{color:var(--text);background:radial-gradient(circle at 85% -10%,rgba(37,99,235,.18),transparent 30%),linear-gradient(180deg,#070b14,#09111f)}
.block-container{padding-top:1.3rem;padding-bottom:3rem;max-width:1320px}h1,h2,h3{color:#f8fafc!important;letter-spacing:-.035em;font-weight:800!important}p,label,.stMarkdown{color:#e2e8f0}
section[data-testid="stSidebar"]{background:linear-gradient(180deg,#111827,#0b1220);border-right:1px solid var(--line)}section[data-testid="stSidebar"] div[role="radiogroup"] label{border:1px solid transparent;border-radius:12px;padding:8px 10px;margin:2px 0}section[data-testid="stSidebar"] div[role="radiogroup"] label:has(input:checked){background:linear-gradient(135deg,rgba(37,99,235,.30),rgba(8,145,178,.16));border-color:rgba(96,165,250,.38)}
div[data-testid="stMetric"]{position:relative;overflow:hidden;background:linear-gradient(145deg,rgba(25,38,61,.98),rgba(13,22,38,.98))!important;border:1px solid rgba(96,165,250,.18);border-radius:20px;padding:20px;min-height:120px;box-shadow:0 16px 38px rgba(0,0,0,.23)}div[data-testid="stMetric"]:before{content:"";position:absolute;left:0;top:0;right:0;height:3px;background:linear-gradient(90deg,#2563eb,#22d3ee)}[data-testid="stMetricLabel"] p{color:#9fb0c7!important;font-weight:650!important}[data-testid="stMetricValue"]{color:#f8fafc!important;font-size:1.9rem!important;font-weight:800!important}
.stTabs [data-baseweb="tab-list"]{gap:7px;background:rgba(15,23,38,.75);border:1px solid var(--line);border-radius:15px;padding:6px;overflow-x:auto}.stTabs [data-baseweb="tab"]{height:42px;border-radius:10px;padding:0 14px;color:#aebdd0}.stTabs [aria-selected="true"]{color:white!important;background:linear-gradient(135deg,#2563eb,#0891b2)!important}.stTabs [data-baseweb="tab-highlight"]{display:none}
div.stButton>button,div.stDownloadButton>button,a[data-testid="stLinkButton"]{color:#f8fafc!important;background:linear-gradient(135deg,#2563eb,#0891b2)!important;border:1px solid rgba(125,211,252,.24)!important;border-radius:12px!important;min-height:44px;font-weight:750!important;box-shadow:0 10px 24px rgba(37,99,235,.22)}div.stButton>button:hover{filter:brightness(1.1);transform:translateY(-1px)}section[data-testid="stSidebar"] div.stButton>button{background:rgba(255,255,255,.04)!important;box-shadow:none}
[data-baseweb="input"]>div,[data-baseweb="select"]>div,[data-baseweb="textarea"]>div,[data-testid="stNumberInput"]>div>div,[data-testid="stDateInput"]>div>div{background:#111b2c!important;border-color:rgba(148,163,184,.2)!important;border-radius:12px!important;color:#f8fafc!important}input,textarea{color:#f8fafc!important}[role="listbox"]{background:#111b2c!important;color:white!important}
[data-testid="stForm"]{background:linear-gradient(145deg,rgba(18,28,46,.94),rgba(11,18,31,.94));border:1px solid var(--line);border-radius:18px;padding:20px}[data-testid="stDataFrame"],[data-testid="stTable"]{border:1px solid var(--line);border-radius:16px;overflow:hidden}
@media(max-width:768px){.block-container{padding:1rem .8rem 2rem}h1{font-size:2rem!important}[data-testid="stMetricValue"]{font-size:1.45rem!important}div[data-testid="stMetric"]{min-height:100px;padding:15px}}

.status-paid{display:inline-block;padding:4px 10px;border-radius:999px;background:rgba(34,197,94,.16);color:#4ade80;font-weight:800;border:1px solid rgba(74,222,128,.3)}
.status-due{display:inline-block;padding:4px 10px;border-radius:999px;background:rgba(239,68,68,.16);color:#f87171;font-weight:800;border:1px solid rgba(248,113,113,.3)}
.status-partial{display:inline-block;padding:4px 10px;border-radius:999px;background:rgba(249,115,22,.16);color:#fb923c;font-weight:800;border:1px solid rgba(251,146,60,.3)}
/* Централна клиентска навигация - компактни бутони 2 по 2 */

.st-key-neighbor_nav_area div[data-testid="stHorizontalBlock"]{
    display:flex!important;
    flex-direction:row!important;
    justify-content:center!important;
    flex-wrap:nowrap!important;
    gap:12px!important;
    margin-bottom:8px!important;
}

/* Ширина на всяка колона с бутон */
.st-key-neighbor_nav_area div[data-testid="stColumn"]{
    flex:0 0 320px!important;
    width:320px!important;
    max-width:320px!important;
    min-width:0!important;
}

/* Контейнер на бутона */
.st-key-neighbor_nav_area div[data-testid="stButton"]{
    width:100%!important;
    margin:0!important;
}

/* Самият бутон */
.st-key-neighbor_nav_area div[data-testid="stButton"] > button{
    width:100%!important;
    min-height:40px!important;
    height:40px!important;
    padding:4px 10px!important;
    border-radius:10px!important;
    font-size:.78rem!important;
    font-weight:750!important;
    line-height:1!important;
    white-space:normal!important;
}

.st-key-neighbor_nav_area div[data-testid="stButton"] > button[kind="secondary"]{
    background:linear-gradient(135deg,#2563eb,#0891b2)!important;
    border:1px solid rgba(125,211,252,.28)!important;
    color:#f8fafc!important;
    box-shadow:0 6px 14px rgba(37,99,235,.18)!important;
}

.st-key-neighbor_nav_area div[data-testid="stButton"] > button[kind="primary"]{
    background:linear-gradient(135deg,#082653,#083b58)!important;
    border:1px solid rgba(96,165,250,.82)!important;
    color:#fff!important;
    box-shadow:
        inset 0 0 0 1px rgba(255,255,255,.09),
        0 8px 18px rgba(2,132,199,.28)!important;
}

.customer-nav-label{
    margin:.15rem 0 .5rem;
    color:#93c5fd;
    font-size:.78rem;
    font-weight:800;
    letter-spacing:.09em;
    text-transform:uppercase;
}

/* Телефон */
@media(max-width:768px){

    .st-key-neighbor_nav_area div[data-testid="stHorizontalBlock"]{
        display:flex!important;
        flex-direction:row!important;
        justify-content:center!important;
        flex-wrap:nowrap!important;
        gap:6px!important;
        margin-bottom:6px!important;
    }

    .st-key-neighbor_nav_area div[data-testid="stColumn"]{
        flex:0 0 calc(50% - 3px)!important;
        width:calc(50% - 3px)!important;
        max-width:calc(50% - 3px)!important;
        min-width:0!important;
    }

    .st-key-neighbor_nav_area div[data-testid="stButton"] > button{
        min-height:38px!important;
        height:38px!important;
        padding:3px 4px!important;
        border-radius:9px!important;
        font-size:.72rem!important;
    }
}

</style>
""", unsafe_allow_html=True)

@st.cache_resource(show_spinner=False)
def initialize_database_schema():
    """Prepare database schema once per Streamlit server process."""
    init_db()
    execute("""
        CREATE TABLE IF NOT EXISTS dashboard_baseline (
            id BIGSERIAL PRIMARY KEY,
            expense_amount NUMERIC(14,2) NOT NULL DEFAULT 7000.00,
            applied_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            applied_by BIGINT REFERENCES users(id),
            note TEXT
        )
    """)

    execute("""
        INSERT INTO dashboard_baseline(expense_amount, applied_at, note)
        SELECT 7000.00, NOW(), 'Начални разходи за текущата година'
        WHERE NOT EXISTS (
            SELECT 1 FROM dashboard_baseline
        )
    """)
    execute("ALTER TABLE payments ADD COLUMN IF NOT EXISTS provider VARCHAR(50)")
    execute("ALTER TABLE payments ADD COLUMN IF NOT EXISTS provider_transaction_id VARCHAR(150)")
    execute("ALTER TABLE payments ADD COLUMN IF NOT EXISTS external_status VARCHAR(40)")
    execute("""CREATE TABLE IF NOT EXISTS balance_baseline (
        id BIGSERIAL PRIMARY KEY, cash_amount NUMERIC(14,2) NOT NULL, bank_amount NUMERIC(14,2) NOT NULL,
        applied_at TIMESTAMPTZ NOT NULL DEFAULT NOW(), applied_by BIGINT REFERENCES users(id), note TEXT)""")
    execute("ALTER TABLE obligations ADD COLUMN IF NOT EXISTS source_currency VARCHAR(3) NOT NULL DEFAULT 'EUR'")
    execute("ALTER TABLE obligations ADD COLUMN IF NOT EXISTS custom_status VARCHAR(30)")
    execute("ALTER TABLE obligations ADD COLUMN IF NOT EXISTS original_amount NUMERIC(12,2)")
    execute("ALTER TABLE obligations ADD COLUMN IF NOT EXISTS adjustment_reason TEXT")
    execute("ALTER TABLE obligations ADD COLUMN IF NOT EXISTS adjusted_by BIGINT REFERENCES users(id)")
    execute("ALTER TABLE obligations ADD COLUMN IF NOT EXISTS adjusted_at TIMESTAMPTZ")
    execute("ALTER TABLE payments ADD COLUMN IF NOT EXISTS source_currency VARCHAR(3) NOT NULL DEFAULT 'EUR'")
    execute("ALTER TABLE apartments ADD COLUMN IF NOT EXISTS monthly_entrance_fee NUMERIC(12,2)")
    execute("ALTER TABLE apartments ADD COLUMN IF NOT EXISTS admin_note TEXT")
    execute("""CREATE TABLE IF NOT EXISTS obligation_adjustments (
        id BIGSERIAL PRIMARY KEY, obligation_id BIGINT NOT NULL REFERENCES obligations(id) ON DELETE CASCADE,
        old_amount NUMERIC(12,2) NOT NULL, new_amount NUMERIC(12,2) NOT NULL, action VARCHAR(30) NOT NULL,
        reason TEXT NOT NULL, changed_by BIGINT REFERENCES users(id), changed_at TIMESTAMPTZ NOT NULL DEFAULT NOW())""")
    execute("""CREATE TABLE IF NOT EXISTS technical_issues (
        id BIGSERIAL PRIMARY KEY, apartment_id BIGINT NOT NULL REFERENCES apartments(id),
        created_by BIGINT REFERENCES users(id), title VARCHAR(200) NOT NULL, description TEXT NOT NULL,
        category VARCHAR(60) NOT NULL, status VARCHAR(30) NOT NULL DEFAULT 'Нова заявка',
        image_name TEXT, image_data BYTEA, image_mime VARCHAR(100),
        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(), updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW())""")
    execute("""CREATE TABLE IF NOT EXISTS documents (
        id BIGSERIAL PRIMARY KEY, title VARCHAR(220) NOT NULL, category VARCHAR(80) NOT NULL,
        description TEXT, file_name TEXT NOT NULL, file_data BYTEA NOT NULL, file_mime VARCHAR(140),
        archived BOOLEAN NOT NULL DEFAULT FALSE, is_deleted BOOLEAN NOT NULL DEFAULT FALSE,
        created_by BIGINT REFERENCES users(id), created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
        archived_at TIMESTAMPTZ, deleted_at TIMESTAMPTZ)""")
    execute("""CREATE TABLE IF NOT EXISTS parking_spaces (
        id BIGSERIAL PRIMARY KEY, space_number VARCHAR(50) UNIQUE,
        apartment_id BIGINT REFERENCES apartments(id), notes TEXT,
        active BOOLEAN NOT NULL DEFAULT TRUE, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW())""")
    execute("""CREATE TABLE IF NOT EXISTS polls (
        id BIGSERIAL PRIMARY KEY, title VARCHAR(220) NOT NULL, question TEXT NOT NULL, description TEXT,
        starts_at TIMESTAMPTZ NOT NULL, ends_at TIMESTAMPTZ NOT NULL,
        active BOOLEAN NOT NULL DEFAULT TRUE, created_by BIGINT REFERENCES users(id),
        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW())""")
    execute("""CREATE TABLE IF NOT EXISTS poll_options (
        id BIGSERIAL PRIMARY KEY, poll_id BIGINT NOT NULL REFERENCES polls(id) ON DELETE CASCADE,
        option_text VARCHAR(300) NOT NULL, sort_order INTEGER NOT NULL DEFAULT 0)""")
    execute("""CREATE TABLE IF NOT EXISTS poll_attachments (
        id BIGSERIAL PRIMARY KEY, poll_id BIGINT NOT NULL REFERENCES polls(id) ON DELETE CASCADE,
        file_name TEXT NOT NULL, file_data BYTEA NOT NULL, file_mime VARCHAR(140),
        is_image BOOLEAN NOT NULL DEFAULT FALSE, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW())""")
    execute("""CREATE TABLE IF NOT EXISTS poll_votes (
        id BIGSERIAL PRIMARY KEY, poll_id BIGINT NOT NULL REFERENCES polls(id) ON DELETE CASCADE,
        option_id BIGINT NOT NULL REFERENCES poll_options(id) ON DELETE CASCADE,
        apartment_id BIGINT NOT NULL REFERENCES apartments(id), user_id BIGINT REFERENCES users(id),
        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(), UNIQUE(poll_id,apartment_id))""")
    execute("ALTER TABLE announcements ADD COLUMN IF NOT EXISTS file_name TEXT")
    execute("ALTER TABLE announcements ADD COLUMN IF NOT EXISTS file_data BYTEA")
    execute("ALTER TABLE announcements ADD COLUMN IF NOT EXISTS file_mime VARCHAR(140)")
    execute("CREATE INDEX IF NOT EXISTS idx_technical_issues_status ON technical_issues(status,created_at DESC)")
    execute("ALTER TABLE documents ADD COLUMN IF NOT EXISTS archived BOOLEAN NOT NULL DEFAULT FALSE")
    execute("ALTER TABLE documents ADD COLUMN IF NOT EXISTS is_deleted BOOLEAN NOT NULL DEFAULT FALSE")
    execute("ALTER TABLE documents ADD COLUMN IF NOT EXISTS archived_at TIMESTAMPTZ")
    execute("ALTER TABLE documents ADD COLUMN IF NOT EXISTS deleted_at TIMESTAMPTZ")
    execute("CREATE INDEX IF NOT EXISTS idx_documents_archive ON documents(archived,is_deleted,created_at DESC)")
    execute("ALTER TABLE poll_votes ADD COLUMN IF NOT EXISTS proof_type VARCHAR(20) NOT NULL DEFAULT 'none'")
    execute("ALTER TABLE poll_votes ADD COLUMN IF NOT EXISTS video_name TEXT")
    execute("ALTER TABLE poll_votes ADD COLUMN IF NOT EXISTS video_data BYTEA")
    execute("ALTER TABLE poll_votes ADD COLUMN IF NOT EXISTS video_mime VARCHAR(100)")
    execute("ALTER TABLE poll_votes ADD COLUMN IF NOT EXISTS paper_signed BOOLEAN NOT NULL DEFAULT FALSE")
    execute("ALTER TABLE poll_votes ADD COLUMN IF NOT EXISTS paper_signed_at TIMESTAMPTZ")
    execute("ALTER TABLE poll_votes ADD COLUMN IF NOT EXISTS paper_signed_by BIGINT REFERENCES users(id)")
    execute("CREATE INDEX IF NOT EXISTS idx_poll_votes_proof ON poll_votes(poll_id,proof_type,paper_signed)")
    execute("CREATE INDEX IF NOT EXISTS idx_polls_period ON polls(starts_at,ends_at)")
    execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS notifications_seen_at TIMESTAMPTZ")
    return True


try:
    initialize_database_schema()
except Exception as exc:
    st.error(str(exc))
    st.stop()


ADMIN_NOTIFICATION_EMAIL = "et.sob.p.vladigerov.74@gmail.com"


def send_admin_email(subject, body):
    """Send an admin alert through SMTP configured in Render environment variables."""
    smtp_host = os.getenv("SMTP_HOST", "smtp.gmail.com").strip()
    smtp_port = int(os.getenv("SMTP_PORT", "587"))
    smtp_user = os.getenv("SMTP_USER", "").strip()
    smtp_password = os.getenv("SMTP_PASSWORD", "").strip()
    smtp_from = os.getenv("SMTP_FROM", smtp_user).strip()

    if not smtp_user or not smtp_password or not smtp_from:
        return False, "Email известието не е изпратено: липсват SMTP_USER/SMTP_PASSWORD в Render."

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = smtp_from
    message["To"] = ADMIN_NOTIFICATION_EMAIL
    message.set_content(body)

    try:
        with smtplib.SMTP(smtp_host, smtp_port, timeout=20) as server:
            server.starttls()
            server.login(smtp_user, smtp_password)
            server.send_message(message)
        return True, ""
    except Exception as exc:
        return False, f"Email известието не беше изпратено: {exc}"


def queue_email_result(subject, body):
    sent, error = send_admin_email(subject, body)
    if not sent:
        st.session_state["email_notification_warning"] = error


def show_email_warning():
    warning = st.session_state.pop("email_notification_warning", None)
    if warning:
        st.warning(warning)


def neighbor_notification_items(user):
    seen_at = fetch_one(
        "SELECT notifications_seen_at FROM users WHERE id=%s",
        (user["id"],)
    )
    seen_at = seen_at.get("notifications_seen_at") if seen_at else None

    params = []
    time_filter = ""
    if seen_at:
        time_filter = " AND created_at > %s"
        params.append(seen_at)

    announcements = fetch_all(f"""
        SELECT id, 'announcement' AS kind, title, created_at
        FROM announcements
        WHERE active = TRUE
          AND (starts_on IS NULL OR starts_on <= CURRENT_DATE)
          AND (ends_on IS NULL OR ends_on >= CURRENT_DATE)
          {time_filter}
    """, tuple(params))

    polls = fetch_all(f"""
        SELECT id, 'poll' AS kind, title, created_at
        FROM polls
        WHERE active = TRUE
          AND starts_at <= NOW()
          AND ends_at >= NOW()
          {time_filter}
    """, tuple(params))

    items = [dict(row) for row in announcements] + [dict(row) for row in polls]
    items.sort(key=lambda row: row["created_at"], reverse=True)
    return items


def neighbor_notification_bell(user):
    items = neighbor_notification_items(user)
    label = f"🔔 {len(items)}" if items else "🔔"

    with st.popover(label, use_container_width=False):
        st.markdown("### Известия")
        if not items:
            st.caption("Няма нови обяви или гласувания.")
        else:
            for item in items:
                icon = "📢" if item["kind"] == "announcement" else "🗳️"
                section = "Обяви" if item["kind"] == "announcement" else "Електронно гласуване"
                if st.button(
                    f"{icon} {item['title']}",
                    key=f"notification_{item['kind']}_{item['id']}",
                    use_container_width=True
                ):
                    execute(
                        "UPDATE users SET notifications_seen_at=NOW() WHERE id=%s",
                        (user["id"],)
                    )
                    _set_neighbor_page(section)
                    rerun()

            if st.button("Маркирай всички като прочетени", key="mark_notifications_read", use_container_width=True):
                execute(
                    "UPDATE users SET notifications_seen_at=NOW() WHERE id=%s",
                    (user["id"],)
                )
                rerun()


def fmt(v): return f"{Decimal(v or 0):,.2f} {CURRENCY}".replace(","," ")
def rerun(): st.rerun()
def method_bg(v): return {"cash":"В брой","bank":"По банка","card_softpos":"Карта / SoftPOS"}.get(v,v)
def payment_status_bg(v): return {"confirmed":"🟢 Платено","pending":"🔴 Неплатено / чака потвърждение","cancelled":"🔴 Отказано","failed":"🔴 Неуспешно"}.get(v,v)
def kind_bg(v): return {"entrance":"Такса вход","fro":"ФРО","old_debt":"Старо задължение"}.get(v,v)
YEARS=list(range(2024,2041))

def login_page():
    banner_candidates = [
        Path("assets/Twilight Luxury Residences with Crescent Moon.jpg"),
        Path("Twilight Luxury Residences with Crescent Moon.jpg"),
    ]
    banner = next((path for path in banner_candidates if path.exists()), None)
    if banner:
        st.image(str(banner), use_container_width=True)
    st.markdown("## Управление на етажна собственост")
    a,b=st.tabs(["Вход","Регистрация за справка"])
    with a:
        with st.form("login"):
            email=st.text_input("Имейл").strip().lower(); pw=st.text_input("Парола",type="password")
            go=st.form_submit_button("Вход",use_container_width=True)
        if go:
            u=fetch_one("SELECT u.*,a.apartment_number FROM users u LEFT JOIN apartments a ON a.id=u.apartment_id WHERE lower(u.email)=lower(%s)",(email,))
            if u and verify_password(pw,u["password_hash"]):
                if u["account_status"]!="approved": st.warning("Профилът очаква одобрение.")
                else: login_user(dict(u)); rerun()
            else: st.error("Невалиден имейл или парола.")
    with b:
        apts=fetch_all("SELECT id,apartment_number FROM apartments WHERE active ORDER BY apartment_number")
        amap={str(x["apartment_number"]):x["id"] for x in apts}
        with st.form("register"):
            apt=st.selectbox("Апартамент / обект",list(amap) or ["Няма обекти"])
            c1,c2=st.columns(2); first=c1.text_input("Име"); last=c2.text_input("Фамилия")
            email=st.text_input("Имейл").strip().lower(); phone=st.text_input("Телефон")
            pw=st.text_input("Парола",type="password"); pw2=st.text_input("Повтори паролата",type="password")
            reg=st.form_submit_button("Изпрати регистрация",use_container_width=True)
        if reg:
            if not amap or not all([first,last,email,pw]) or pw!=pw2 or len(pw)<8: st.error("Провери данните и паролата.")
            else:
                try:
                    execute("INSERT INTO users(email,password_hash,first_name,last_name,phone,role,apartment_id,account_status) VALUES(%s,%s,%s,%s,%s,'neighbor',%s,'pending')",(email,hash_password(pw),first,last,phone,amap[apt]))
                    queue_email_result(
                        f"Building.Manager: нова регистрация за обект {apt}",
                        f"Нова регистрация очаква одобрение.\n\nИме: {first} {last}\nОбект: {apt}\nИмейл: {email}\nТелефон: {phone or '-'}"
                    )
                    st.success("Регистрацията е изпратена.")
                except Exception: st.error("Имейлът вече е използван.")

def payment_rows(apartment_id, year):
    create_year_obligations(apartment_id, int(year))

    apartment = fetch_one("""
        SELECT adult_residents, monthly_entrance_fee
        FROM apartments
        WHERE id = %s
    """, (apartment_id,))

    if apartment:
        entrance_fee = (
            Decimal(apartment["monthly_entrance_fee"])
            if apartment.get("monthly_entrance_fee") is not None
            else Decimal(apartment["adult_residents"] or 0) * Decimal("6.00")
        )

        execute("""
            UPDATE obligations o
            SET amount = %s
            WHERE o.apartment_id = %s
              AND o.year = %s
              AND o.fee_type = 'entrance'
              AND o.custom_status IS DISTINCT FROM 'force_majeure'
              AND COALESCE(
                    (
                        SELECT SUM(pa.amount)
                        FROM payment_allocations pa
                        JOIN payments p ON p.id = pa.payment_id
                        WHERE pa.obligation_id = o.id
                          AND p.status = 'confirmed'
                    ),
                    0
                  ) = 0
        """, (entrance_fee, apartment_id, int(year)))

    return obligation_rows(apartment_id, int(year))

def apartment_ledger(
    apartment_id,
    year,
    editable=False,
    key="ledger",
    rows=None
):
    if rows is None:
        rows = payment_rows(apartment_id, year)

    data = []
    for r in rows:
        paid=Decimal(r["paid"]); left=max(Decimal("0"),Decimal(r["amount"])-paid)
        custom_status = str(r.get("custom_status") or "")
        if custom_status == "force_majeure": display_status = "🔵 Друг / форсмажор"
        elif left <= 0: display_status = "🟢 Платено"
        elif paid > 0: display_status = "🟠 Частично"
        else: display_status = "🔴 Неплатено"
        data.append({"Плати":False,"Вид":kind_bg(r["fee_type"]),"Месец":MONTHS.get(r["month"],"Еднократно"),"Година":r["year"],"Валута":r.get("source_currency") or ("BGN" if int(r["year"])<2026 else "EUR"),"Начислено":float(r["amount"]),"Платено":float(paid),"Остатък":float(left),"Статус":display_status,"_id":r["id"]})
    df=pd.DataFrame(data)
    if df.empty: st.info("Няма задължения."); return [],Decimal("0")
    base=df.drop(columns=["_id"])
    if editable:
        view=st.data_editor(base,use_container_width=True,hide_index=True,disabled=["Вид","Месец","Година","Валута","Начислено","Платено","Остатък","Статус"],column_config={"Плати":st.column_config.CheckboxColumn("Плати"),"Начислено":st.column_config.NumberColumn(format="%.2f"),"Платено":st.column_config.NumberColumn(format="%.2f"),"Остатък":st.column_config.NumberColumn(format="%.2f")},key=key)
    else:
        view=base.drop(columns=["Плати"]); st.dataframe(view,use_container_width=True,hide_index=True); return [],sum(Decimal(str(x)) for x in df["Остатък"])
    selected=[]
    for i,row in view.iterrows():
        if bool(row["Плати"]) and Decimal(str(row["Остатък"]))>0: selected.append({"id":int(df.iloc[i]["_id"]),"amount":Decimal(str(row["Остатък"])),"label":f"{row['Вид']} {row['Месец']} {row['Година']}"})
    return selected,sum((x["amount"] for x in selected),Decimal("0"))


def selected_currency_for_year(year):
    return "BGN" if int(year) < 2026 else "EUR"


def record_payment_with_currency(apartment_id, selected, method, user_id, note, currency, status="confirmed"):
    number=create_payment(apartment_id,selected,method,user_id,status,note)
    execute("UPDATE payments SET source_currency=%s WHERE payment_number=%s",(currency,number))
    return number


def force_majeure_adjustment(user, apartment_id, year, rows):
    if user.get("role") != "manager": return
    st.divider(); st.subheader("⚙️ Друг статус / форсмажорна корекция")
    st.caption("Променя само избраното задължение. Не създава приход или разход и не променя касовото салдо.")
    row_map={}
    for row in rows:
        paid=Decimal(row["paid"]); amount=Decimal(row["amount"]); remaining=max(Decimal("0"),amount-paid)
        state="ПЛАТЕНО" if remaining==0 else "ЧАСТИЧНО" if paid>0 else "НЕПЛАТЕНО"
        adjusted=" · ДРУГ" if row.get("custom_status")=="force_majeure" else ""
        label=f"{state}{adjusted} · {kind_bg(row['fee_type'])} · {MONTHS.get(row['month'],'Еднократно')} {row['year']} · {amount:.2f} {row.get('source_currency') or 'EUR'}"
        row_map[label]={"row":row,"paid":paid,"amount":amount}
    if not row_map: st.info("Няма задължения за корекция."); return
    selected_label=st.selectbox("Избери задължение",list(row_map),key=f"fm_obligation_{apartment_id}_{year}")
    item=row_map[selected_label]; row=item["row"]; paid=item["paid"]; current=item["amount"]
    is_adjusted=row.get("custom_status")=="force_majeure"
    actions=["Без промяна","Друг / форсмажор"]
    if is_adjusted and row.get("original_amount") is not None: actions.append("Върни първоначалната сума")
    action=st.selectbox("Действие",actions,key=f"fm_action_{row['id']}")
    if action=="Без промяна": return
    if action=="Върни първоначалната сума":
        original=Decimal(row["original_amount"])
        if original < paid:
            st.error(f"Първоначалната сума {original:.2f} е под вече платеното {paid:.2f}. Възстановяването е блокирано."); return
        reason=st.text_input("Причина за възстановяване",key=f"restore_reason_{row['id']}")
        confirm=st.checkbox("Потвърждавам възстановяването",key=f"restore_confirm_{row['id']}")
        if st.button("Върни първоначалната сума",disabled=not(reason.strip() and confirm),key=f"restore_btn_{row['id']}"):
            execute("INSERT INTO obligation_adjustments(obligation_id,old_amount,new_amount,action,reason,changed_by) VALUES(%s,%s,%s,'restore',%s,%s)",(row["id"],current,original,reason.strip(),user["id"]))
            execute("UPDATE obligations SET amount=%s,custom_status=NULL,adjustment_reason=NULL,adjusted_by=%s,adjusted_at=NOW() WHERE id=%s",(original,user["id"],row["id"]))
            audit(user["id"],"restore_obligation_amount","obligation",row["id"],f"{current}->{original}; {reason}"); st.success("Първоначалната сума е възстановена."); rerun()
        return
    st.warning("Използвай само при решение, освобождаване, индивидуална договорка или допълнително начисление.")
    with st.form(f"fm_form_{row['id']}"):
        new_amount=st.number_input("Нова начислена сума",min_value=float(paid),max_value=100000.0,value=float(current),step=0.01,format="%.2f")
        reason=st.text_area("Задължително основание",placeholder="Решение, форсмажор, освобождаване или допълнително начисление")
        confirm=st.checkbox("Потвърждавам корекцията")
        save=st.form_submit_button("Запази корекцията",use_container_width=True)
    if save:
        corrected=Decimal(str(new_amount)).quantize(Decimal("0.01"))
        if not reason.strip(): st.error("Въведи основание."); return
        if not confirm: st.error("Потвърди корекцията."); return
        original=Decimal(row["original_amount"]) if row.get("original_amount") is not None else current
        execute("INSERT INTO obligation_adjustments(obligation_id,old_amount,new_amount,action,reason,changed_by) VALUES(%s,%s,%s,'force_majeure',%s,%s)",(row["id"],current,corrected,reason.strip(),user["id"]))
        execute("""UPDATE obligations SET amount=%s,original_amount=COALESCE(original_amount,%s),custom_status='force_majeure',adjustment_reason=%s,adjusted_by=%s,adjusted_at=NOW() WHERE id=%s""",(corrected,original,reason.strip(),user["id"],row["id"]))
        audit(user["id"],"force_majeure_adjustment","obligation",row["id"],f"{current}->{corrected}; {reason}")
        st.success(f"Корекцията е записана: {current:.2f} → {corrected:.2f}."); rerun()

def cashier_page(user):
    st.title("💳 Каса")
    apts=fetch_all("SELECT id,apartment_number,property_type,adult_residents,annual_fro,monthly_entrance_fee FROM apartments WHERE active ORDER BY apartment_number")
    amap={f"{a['apartment_number']} · {a['property_type']}":a for a in apts}
    if not amap: st.warning("Няма добавени обекти."); return
    selected_name=st.selectbox("Апартамент / обект",list(amap)); apt=amap[selected_name]
    default_year=YEARS.index(date.today().year) if date.today().year in YEARS else 0
    year=st.selectbox("Година",YEARS,index=default_year)
    rows=payment_rows(apt["id"],year); paid=sum(Decimal(x["paid"]) for x in rows); due=sum(max(Decimal("0"),Decimal(x["amount"])-Decimal(x["paid"])) for x in rows)
    c1,c2,c3,c4=st.columns(4); c1.metric("Обект",apt["apartment_number"]); c2.metric("Пълнолетни",apt["adult_residents"]); c3.metric("Платено",fmt(paid)); c4.metric("Остава",fmt(due))
    st.subheader("Финансов картон")
    selected,total=apartment_ledger(apt["id"],year,True,f"ledger_{apt['id']}_{year}",rows=rows)
    force_majeure_adjustment(user,apt["id"],year,rows)
    currency=selected_currency_for_year(year)
    st.divider(); c1,c2=st.columns([2,1]); c1.write(f"**Избрани позиции: {len(selected)}**"); c2.metric("Обща сума",f"{total:,.2f} {currency}".replace(","," "))
    note=st.text_input("Бележка / номер на платежно")
    b1,b2,b3=st.columns(3)
    if b1.button("💵 В брой",use_container_width=True,disabled=not selected): record_payment_with_currency(apt["id"],selected,"cash",user["id"],note,currency); st.success("Плащането е записано."); rerun()
    if b2.button("🏦 По банка",use_container_width=True,disabled=not selected): record_payment_with_currency(apt["id"],selected,"bank",user["id"],note,currency); st.success("Плащането е записано."); rerun()
    if b3.button("💳 Карта / NFC",use_container_width=True,disabled=not selected):
        if currency != "EUR": st.error("Картово плащане е разрешено само за EUR задължения."); return
        num=record_payment_with_currency(apt["id"],selected,"card_softpos",user["id"],note or "SoftPOS",currency,"pending")
        p=fetch_one("SELECT id FROM payments WHERE payment_number=%s",(num,)); execute("UPDATE payments SET provider='mypos_glass',external_status='initiated' WHERE id=%s",(p["id"],))
        st.session_state.softpos={"id":p["id"],"number":num,"amount":str(total)}; rerun()
    pending=st.session_state.get("softpos")
    if pending:
        st.warning(f"SoftPOS операция **{pending['number']}**, сума **{fmt(pending['amount'])}**. Потвърди само след успешно плащане в myPOS Glass.")
        template=os.getenv("MYPOS_GLASS_DEEP_LINK_TEMPLATE","")
        if template:
            url=template.replace("{amount}",quote(pending["amount"])).replace("{reference}",quote(pending["number"])).replace("{currency}","EUR")
            st.link_button("Отвори myPOS Glass",url,use_container_width=True)
        else: st.info("Стартирай сумата ръчно в myPOS Glass. Автоматичният бутон се активира след добавяне на одобрения deep-link шаблон в Render.")
        ref=st.text_input("Референция от успешната картова транзакция")
        x,y=st.columns(2)
        if x.button("✅ Потвърди успешно плащане",disabled=not ref.strip(),use_container_width=True):
            execute("UPDATE payments SET status='confirmed',payment_date=CURRENT_DATE,confirmed_by=%s,confirmed_at=NOW(),provider_transaction_id=%s,external_status='approved' WHERE id=%s AND status='pending'",(user["id"],ref.strip(),pending["id"])); audit(user["id"],"confirm_softpos","payment",pending["id"],ref); st.session_state.pop("softpos",None); rerun()
        if y.button("Откажи операцията",use_container_width=True): execute("UPDATE payments SET status='cancelled',external_status='cancelled' WHERE id=%s",(pending["id"],)); st.session_state.pop("softpos",None); rerun()
    st.subheader("Последни плащания")
    hist=fetch_all("SELECT payment_date AS дата,payment_number AS номер,amount AS сума,payment_method AS метод,status AS статус,note AS бележка FROM payments WHERE apartment_id=%s ORDER BY created_at DESC LIMIT 20",(apt["id"],))
    st.dataframe([{**dict(x),"метод":method_bg(x["метод"]),"статус":payment_status_bg(x["статус"])} for x in hist],use_container_width=True,hide_index=True)

def report_page(user):
    st.title("📊 Отчет")
    base=fetch_one("SELECT * FROM balance_baseline ORDER BY applied_at DESC LIMIT 1")
    if not base:
        st.warning("Началните салда още не са приложени от Админ панел → Исторически статус.")
        base={"cash_amount":Decimal("0"),"bank_amount":Decimal("0"),"applied_at":date.today()}
    cash_in=fetch_one("SELECT COALESCE(SUM(amount),0) total FROM payments WHERE status='confirmed' AND payment_method='cash' AND source_currency='EUR' AND created_at>%s",(base["applied_at"],))["total"]
    bank_in=fetch_one("SELECT COALESCE(SUM(amount),0) total FROM payments WHERE status='confirmed' AND payment_method IN ('bank','card_softpos') AND source_currency='EUR' AND created_at>%s",(base["applied_at"],))["total"]
    cash_out=fetch_one("SELECT COALESCE(SUM(amount),0) total FROM expenses WHERE source='cash' AND created_at>%s",(base["applied_at"],))["total"]
    bank_out=fetch_one("SELECT COALESCE(SUM(amount),0) total FROM expenses WHERE source='bank' AND created_at>%s",(base["applied_at"],))["total"]
    cash_balance=Decimal(base["cash_amount"])+Decimal(cash_in)-Decimal(cash_out)
    bank_balance=Decimal(base["bank_amount"])+Decimal(bank_in)-Decimal(bank_out)
    c1,c2,c3=st.columns(3); c1.metric("Каса в брой",fmt(cash_balance)); c2.metric("Каса банка",fmt(bank_balance)); c3.metric("Всичко",fmt(cash_balance+bank_balance))
    st.caption(f"Начално салдо: в брой {fmt(base['cash_amount'])}, банка {fmt(base['bank_amount'])}. След него отчетът добавя новите постъпления и изважда новите разходи.")
    st.subheader("Поток на постъпления и разходи")
    payments=fetch_all("""SELECT p.created_at AS дата,'Приход' AS поток,a.apartment_number AS обект,p.payment_number AS документ,
        p.amount AS сума,p.source_currency AS валута,CASE p.payment_method WHEN 'cash' THEN 'В брой' WHEN 'bank' THEN 'По банка' ELSE 'Карта / SoftPOS' END AS начин,
        CASE p.status WHEN 'confirmed' THEN '🟢 Платено' WHEN 'pending' THEN '🔴 Неплатено / чака' ELSE '🔴 Отказано' END AS статус
        FROM payments p JOIN apartments a ON a.id=p.apartment_id WHERE p.created_at>%s""",(base["applied_at"],))
    expenses=fetch_all("""SELECT created_at AS дата,'Разход' AS поток,category AS обект,COALESCE(description,'') AS документ,
        -amount AS сума,'EUR' AS валута,CASE source WHEN 'cash' THEN 'В брой' ELSE 'По банка' END AS начин,'Разход' AS статус
        FROM expenses WHERE created_at>%s""",(base["applied_at"],))
    flow=[dict(x) for x in payments]+[dict(x) for x in expenses]; flow.sort(key=lambda x:x["дата"],reverse=True)
    st.dataframe(flow,use_container_width=True,hide_index=True)
    if flow:
        export_df=pd.DataFrame(flow).copy()
        if "дата" in export_df.columns:
            export_df["дата"]=pd.to_datetime(export_df["дата"],errors="coerce",utc=True).dt.tz_localize(None)
        out=io.BytesIO()
        export_df.to_excel(out,index=False,engine="openpyxl")
        st.download_button("Изтегли отчет",data=out.getvalue(),file_name=f"building_report_{date.today()}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)

def expenses_page(user):
    st.title("Разходи")
    # Наличност каса за съседите

    base = fetch_one("""
        SELECT *
        FROM balance_baseline
        ORDER BY applied_at DESC
        LIMIT 1
    """)

    if base:

        cash_in = fetch_one("""
            SELECT COALESCE(SUM(amount),0) total
            FROM payments
            WHERE status='confirmed'
              AND payment_method='cash'
              AND source_currency='EUR'
              AND created_at > %s
        """,(base["applied_at"],))["total"]

        bank_in = fetch_one("""
            SELECT COALESCE(SUM(amount),0) total
            FROM payments
            WHERE status='confirmed'
              AND payment_method IN ('bank','card_softpos')
              AND source_currency='EUR'
              AND created_at > %s
        """,(base["applied_at"],))["total"]

        cash_out = fetch_one("""
            SELECT COALESCE(SUM(amount),0) total
            FROM expenses
            WHERE source='cash'
              AND created_at > %s
        """,(base["applied_at"],))["total"]

        bank_out = fetch_one("""
            SELECT COALESCE(SUM(amount),0) total
            FROM expenses
            WHERE source='bank'
              AND created_at > %s
        """,(base["applied_at"],))["total"]

        cash_balance = (
            Decimal(base["cash_amount"])
            + Decimal(cash_in)
            - Decimal(cash_out)
        )

        bank_balance = (
            Decimal(base["bank_amount"])
            + Decimal(bank_in)
            - Decimal(bank_out)
        )

        total_balance = cash_balance + bank_balance

        st.metric(
            "💶 Наличност Каса",
            fmt(total_balance)
        )

        st.caption(
            "Актуална наличност на етажната собственост."
        )

        st.divider()    
    if user["role"] in ("manager","cashier"):
        cats=[x["category"] for x in fetch_all("SELECT DISTINCT category FROM expenses ORDER BY category")]
        defaults=["Електроенергия","Асансьор","Почистване","Интернет","Банкови такси","Домоуправител","Касиер","Ремонт","Консумативи","Гараж","Други"]
        with st.form("expense"):
            d=st.date_input("Дата",date.today()); cat=st.selectbox("Категория",list(dict.fromkeys(defaults+cats))); amount=st.number_input("Сума (€)",0.0,100000.0,0.0,.01); source=st.radio("Начин",["В брой","По банка"],horizontal=True); desc=st.text_area("Описание"); save=st.form_submit_button("Добави разход")
        if save and amount>0: execute("INSERT INTO expenses(expense_date,category,amount,source,description,created_by) VALUES(%s,%s,%s,%s,%s,%s)",(d,cat,amount,"cash" if source=="В брой" else "bank",desc,user["id"])); rerun()
    rows=fetch_all("SELECT expense_date AS дата,category AS категория,amount AS сума,CASE source WHEN 'cash' THEN 'В брой' ELSE 'По банка' END AS начин,description AS описание FROM expenses ORDER BY expense_date DESC,id DESC")
    st.dataframe(rows,use_container_width=True,hide_index=True)
def _set_neighbor_page(page_name):
    st.session_state.neighbor_page = page_name


def neighbor_main_navigation():
    """Точно два native Streamlit бутона на ред без URL презареждане."""
    if "neighbor_page" not in st.session_state:
        st.session_state.neighbor_page = "Моите задължения"

    nav_items = [
        ("Моите задължения", "🏠 Задължения"),
        ("Архив", "🗂️ Архив"),
        ("Разходи", "💶 Разходи"),
        ("Технически проблем", "🛠️ Проблем"),
        ("Обяви", "📢 Обяви"),
        ("Документи", "📁 Документи"),
        ("Електронно гласуване", "🗳️ Вот"),
        ("Паркоместа", "🅿️ Паркинг"),
    ]

    st.markdown('<div class="customer-nav-label">Бърз достъп</div>', unsafe_allow_html=True)

    with st.container(key="neighbor_nav_area"):
        for row_start in range(0, len(nav_items), 2):
            left, right = st.columns(2, gap="small")
            pair = nav_items[row_start:row_start + 2]
            for column, item in zip((left, right), pair):
                page_name, label = item
                with column:
                    st.button(
                        label,
                        key=f"neighbor_nav_{page_name}",
                        type="primary" if st.session_state.neighbor_page == page_name else "secondary",
                        use_container_width=True,
                        on_click=_set_neighbor_page,
                        args=(page_name,),
                    )

    st.divider()

def neighbor_page(user):
    st.title("Моите задължения")
    st.caption("🔄 За по-добра визуализация на таблицата завъртете телефона хоризонтално.")
    if not user.get("apartment_id"): st.info("Профилът не е свързан с обект."); return
    default_year=YEARS.index(date.today().year) if date.today().year in YEARS else 0
    year=st.selectbox("Година",YEARS,index=default_year); apartment_ledger(user["apartment_id"],year,False)

def neighbor_archive(user):
    st.title("Архив")
    rows=fetch_all("SELECT payment_date AS дата,payment_number AS номер,amount AS сума,payment_method AS метод,status AS статус,note AS бележка FROM payments WHERE apartment_id=%s AND status='confirmed' ORDER BY payment_date DESC",(user["apartment_id"],)) if user.get("apartment_id") else []
    st.dataframe([{**dict(x),"метод":method_bg(x["метод"]),"статус":payment_status_bg(x["статус"])} for x in rows],use_container_width=True,hide_index=True)



MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_VOTE_VIDEO_BYTES = 12 * 1024 * 1024


def safe_vote_video(upload):
    """Validate a short vote-confirmation video without re-encoding it."""
    if not upload:
        return None, None, None
    data = upload.getvalue()
    mime = upload.type or "application/octet-stream"
    allowed = ("video/mp4", "video/webm", "video/quicktime")
    if mime not in allowed:
        raise ValueError("Разрешени са MP4, WEBM и MOV видео файлове.")
    if len(data) > MAX_VOTE_VIDEO_BYTES:
        raise ValueError("Видеото е по-голямо от 12 MB. Запишете кратък клип до 10 секунди.")
    return upload.name, data, mime


def safe_upload(upload, allowed_prefixes=None):
    if not upload: return None, None, None
    data=upload.getvalue()
    if len(data)>MAX_UPLOAD_BYTES: raise ValueError("Файлът е по-голям от 15 MB.")
    mime=upload.type or "application/octet-stream"
    if allowed_prefixes and not any(mime.startswith(prefix) for prefix in allowed_prefixes):
        raise ValueError("Неподдържан тип файл.")
    return upload.name,data,mime


def safe_image(upload):
    if not upload: return None,None,None
    raw=upload.getvalue()
    if len(raw)>MAX_UPLOAD_BYTES: raise ValueError("Снимката е по-голяма от 15 MB.")
    im=Image.open(io.BytesIO(raw)).convert("RGB")
    im.thumbnail((1600,1200))
    out=io.BytesIO(); im.save(out,"WEBP",quality=82,method=6)
    return getattr(upload,"name","camera.webp"),out.getvalue(),"image/webp"


def technical_issues_page(user):
    st.title("🛠️ Технически проблем")
    if not user.get("apartment_id"):
        st.info("Профилът не е свързан с апартамент или обект."); return
    use_camera=st.checkbox("Използвай камера",key="issue_use_camera")
    source=st.camera_input("Направете снимка",key="issue_camera") if use_camera else st.file_uploader("Качване от файл",type=["png","jpg","jpeg","webp"],key="issue_file")
    with st.form("technical_issue_form"):
        title=st.text_input("Заглавие")
        description=st.text_area("Описание")
        category=st.selectbox("Категория",["Осветление","Асансьор","Врата/достъп","ВиК","Електро","Почистване","Друг проблем"])
        submit=st.form_submit_button("Изпрати сигнал",use_container_width=True)
    if submit:
        if not title.strip() or not description.strip(): st.error("Заглавието и описанието са задължителни.")
        else:
            try: name,data,mime=safe_image(source)
            except Exception as exc: st.error(str(exc)); return
            execute("""INSERT INTO technical_issues(apartment_id,created_by,title,description,category,image_name,image_data,image_mime)
                VALUES(%s,%s,%s,%s,%s,%s,%s,%s)""",(user["apartment_id"],user["id"],title.strip(),description.strip(),category,name,data,mime))
            apartment = fetch_one("SELECT apartment_number FROM apartments WHERE id=%s", (user["apartment_id"],))
            apartment_number = apartment["apartment_number"] if apartment else str(user["apartment_id"])
            queue_email_result(
                f"Building.Manager: нов технически сигнал от обект {apartment_number}",
                f"Подаден е нов технически сигнал.\n\nОбект: {apartment_number}\nКатегория: {category}\nЗаглавие: {title.strip()}\nОписание: {description.strip()}"
            )
            audit(user["id"],"create_technical_issue","technical_issue",None,title.strip())
            st.success("Сигналът е изпратен."); rerun()
    st.subheader("Моите сигнали")
    issues=fetch_all("SELECT * FROM technical_issues WHERE apartment_id=%s ORDER BY created_at DESC",(user["apartment_id"],))
    if not issues: st.info("Все още няма подадени сигнали.")
    for issue in issues:
        icon={"Нова заявка":"🔴","В процес":"🟠","Решена":"🟢"}.get(issue["status"],"🔵")
        with st.expander(f"{icon} {issue['status']} · {issue['title']} · {issue['created_at']:%d.%m.%Y}"):
            st.write(f"**Категория:** {issue['category']}"); st.write(issue["description"])
            if issue["image_data"]: st.image(bytes(issue["image_data"]),width=520)


def announcements_page():
    st.title("📢 Обяви")
    rows=fetch_all("""SELECT * FROM announcements WHERE active
        AND (starts_on IS NULL OR starts_on<=CURRENT_DATE)
        AND (ends_on IS NULL OR ends_on>=CURRENT_DATE) ORDER BY created_at DESC""")
    if not rows: st.info("Няма активни обяви.")
    for row in rows:
        with st.container(border=True):
            if row["image_data"]: st.image(bytes(row["image_data"]),use_container_width=True)
            st.subheader(row["title"]); st.write(row["body"] or "")
            st.caption(f"Публикувано: {(row['starts_on'] or row['created_at'].date()):%d.%m.%Y}")
            if row.get("file_data"):
                st.download_button("Изтегли прикачения файл",bytes(row["file_data"]),row["file_name"] or "attachment",row["file_mime"] or "application/octet-stream",key=f"ann_download_{row['id']}")


def parking_page():
    st.title("🅿️ Паркоместа")
    st.info("Модулът ще бъде активиран в бъдеща версия.")


def documents_page():
    st.title("📁 Документи")
    active,archive=st.tabs(["Активни документи","Архив"])
    for tab,archived in ((active,False),(archive,True)):
        with tab:
            rows=fetch_all("SELECT * FROM documents WHERE archived=%s AND is_deleted=FALSE ORDER BY created_at DESC",(archived,))
            if not rows: st.info("Няма документи в тази секция.")
            for row in rows:
                with st.expander(f"{row['category']} · {row['title']}"):
                    st.write(row["description"] or ""); st.caption(f"Качен: {row['created_at']:%d.%m.%Y}")
                    st.download_button("Преглед / Изтегляне",bytes(row["file_data"]),row["file_name"],row["file_mime"] or "application/octet-stream",key=f"doc_download_{row['id']}")


def voting_page(user):
    st.title("🗳️ Електронно гласуване")
    if not user.get("apartment_id"):
        st.info("Профилът не е свързан с апартамент или обект.")
        return

    polls=fetch_all("SELECT * FROM polls WHERE active AND starts_at<=NOW() AND ends_at>=NOW() ORDER BY ends_at")
    if not polls:
        st.info("Няма активно гласуване.")

    for poll in polls:
        with st.container(border=True):
            st.subheader(poll["title"])
            st.write(f"**{poll['question']}**")
            st.write(poll["description"] or "")
            st.caption(f"Период: {poll['starts_at']:%d.%m.%Y %H:%M} – {poll['ends_at']:%d.%m.%Y %H:%M}")

            for att in fetch_all("SELECT * FROM poll_attachments WHERE poll_id=%s ORDER BY id",(poll["id"],)):
                if att["is_image"]:
                    st.image(bytes(att["file_data"]),width=600)
                st.download_button(
                    f"Изтегли {att['file_name']}",
                    bytes(att["file_data"]),
                    att["file_name"],
                    att["file_mime"] or "application/octet-stream",
                    key=f"poll_attachment_{att['id']}",
                )

            previous=fetch_one("""SELECT pv.option_id,pv.proof_type,pv.paper_signed,
                (pv.video_data IS NOT NULL) AS has_video
                FROM poll_votes pv WHERE pv.poll_id=%s AND pv.apartment_id=%s""",
                (poll["id"],user["apartment_id"]))
            options=fetch_all("SELECT * FROM poll_options WHERE poll_id=%s ORDER BY sort_order,id",(poll["id"],))
            option_map={o["option_text"]:o["id"] for o in options}

            if previous:
                chosen=next((o["option_text"] for o in options if o["id"]==previous["option_id"]),"")
                proof = "🎥 Видео" if previous.get("has_video") else "✍️ Хартиен подпис" if previous.get("paper_signed") else "Очаква доказателство"
                st.success(f"Апартаментът вече е гласувал: {chosen} · {proof}")
            elif options:
                choice=st.radio("Изберете отговор",list(option_map),key=f"poll_choice_{poll['id']}")
                st.markdown("**Потвърждение на гласа**")
                proof_choice=st.radio(
                    "Изберете начин",
                    ["Кратко видео", "Ще подпиша на хартия"],
                    horizontal=True,
                    key=f"poll_proof_{poll['id']}",
                )

                video=None
                if proof_choice == "Кратко видео":
                    st.caption("Запишете клип до 10 секунди, в който казвате номера на апартамента и избрания отговор.")
                    video=st.file_uploader(
                        "Запиши или избери видео",
                        type=["mp4","webm","mov"],
                        key=f"poll_video_{poll['id']}",
                        help="На телефон бутонът позволява избор на заснет видеоклип според браузъра и устройството.",
                    )
                    if video:
                        st.video(video.getvalue())
                else:
                    st.info("Гласът ще бъде записан, а домоуправителят ще отбележи подписа след подписване на хартиения списък.")

                confirm=st.checkbox(
                    f"Потвърждавам отговора „{choice}“ за този апартамент",
                    key=f"poll_confirm_{poll['id']}",
                )
                disabled = not confirm or (proof_choice == "Кратко видео" and video is None)

                if st.button("Потвърди и подай глас",key=f"poll_vote_{poll['id']}",type="primary",disabled=disabled):
                    try:
                        video_name=video_data=video_mime=None
                        proof_type="paper_pending"
                        if proof_choice == "Кратко видео":
                            video_name,video_data,video_mime=safe_vote_video(video)
                            proof_type="video"
                        execute("""INSERT INTO poll_votes(
                            poll_id,option_id,apartment_id,user_id,proof_type,video_name,video_data,video_mime
                            ) VALUES(%s,%s,%s,%s,%s,%s,%s,%s)""",
                            (poll["id"],option_map[choice],user["apartment_id"],user["id"],proof_type,video_name,video_data,video_mime))
                        apartment = fetch_one("SELECT apartment_number FROM apartments WHERE id=%s", (user["apartment_id"],))
                        apartment_number = apartment["apartment_number"] if apartment else str(user["apartment_id"])
                        queue_email_result(
                            f"Building.Manager: подаден глас от обект {apartment_number}",
                            f"Подаден е нов глас.\n\nГласуване: {poll['title']}\nВъпрос: {poll['question']}\nОбект: {apartment_number}\nОтговор: {choice}\nПотвърждение: {proof_type}"
                        )
                        audit(user["id"],"vote","poll",poll["id"],f"{choice}; proof={proof_type}")
                        st.success("Гласът е записан.")
                        rerun()
                    except ValueError as exc:
                        st.error(str(exc))
                    except Exception:
                        st.warning("Този апартамент вече е гласувал или видеото не можа да бъде записано.")

def admin_technical_tab(user):
    issues=fetch_all("""SELECT ti.*,a.apartment_number FROM technical_issues ti
        JOIN apartments a ON a.id=ti.apartment_id
        ORDER BY CASE ti.status WHEN 'Нова заявка' THEN 1 WHEN 'В процес' THEN 2 ELSE 3 END,ti.created_at DESC""")
    if not issues: st.info("Няма технически сигнали.")
    for issue in issues:
        with st.expander(f"Ап. {issue['apartment_number']} · {issue['status']} · {issue['title']}"):
            st.write(f"**Категория:** {issue['category']}"); st.write(issue["description"])
            if issue["image_data"]: st.image(bytes(issue["image_data"]),width=520)
            statuses=["Нова заявка","В процес","Решена"]
            status=st.selectbox("Статус",statuses,index=statuses.index(issue["status"]) if issue["status"] in statuses else 0,key=f"issue_status_{issue['id']}")
            if st.button("Запази статус",key=f"issue_save_{issue['id']}"):
                execute("UPDATE technical_issues SET status=%s,updated_at=NOW() WHERE id=%s",(status,issue["id"]))
                audit(user["id"],"update_technical_issue","technical_issue",issue["id"],status); rerun()


def admin_announcements_tab(user):
    st.subheader("Нова обява")
    title=st.text_input("Заглавие",key="announcement_title")
    body=st.text_area("Текст",key="announcement_body")
    c1,c2=st.columns(2)
    image=c1.file_uploader("Снимка",type=["png","jpg","jpeg","webp"],key="announcement_image")
    attachment=c2.file_uploader("Файл",key="announcement_file")
    date_col,active_col=st.columns(2)
    published=date_col.date_input("Дата на публикуване",date.today(),key="announcement_date")
    active=active_col.checkbox("Активна",True,key="announcement_active")
    if st.button("Публикувай обявата",type="primary",use_container_width=True):
        if not title.strip(): st.error("Заглавието е задължително.")
        else:
            try:
                _,image_data,image_mime=safe_image(image)
                file_name,file_data,file_mime=safe_upload(attachment)
            except Exception as exc: st.error(str(exc)); return
            execute("""INSERT INTO announcements(title,body,image_data,image_mime,file_name,file_data,file_mime,active,starts_on,created_by)
                VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",(title.strip(),body,image_data,image_mime,file_name,file_data,file_mime,active,published,user["id"]))
            audit(user["id"],"create_announcement","announcement",None,title.strip()); rerun()
    st.subheader("Всички обяви")
    for row in fetch_all("SELECT id,title,active,starts_on,created_at FROM announcements ORDER BY created_at DESC"):
        c1,c2=st.columns([5,1]); c1.write(f"**{row['title']}** · {'Активна' if row['active'] else 'Неактивна'} · {(row['starts_on'] or row['created_at'].date()):%d.%m.%Y}")
        if c2.button("Вкл./Изкл.",key=f"announcement_toggle_{row['id']}"):
            execute("UPDATE announcements SET active=NOT active WHERE id=%s",(row["id"],)); rerun()


def admin_documents_tab(user):
    st.subheader("Качване на документ")
    with st.form("admin_document_form"):
        title=st.text_input("Заглавие")
        category=st.selectbox("Категория",["Протоколи","Бюджети","Договори","Решения на общото събрание","Други документи"])
        description=st.text_area("Описание")
        upload=st.file_uploader("Документ")
        save=st.form_submit_button("Запази документ",use_container_width=True)
    if save:
        if not title.strip() or not upload: st.error("Заглавието и документът са задължителни.")
        else:
            try: name,data,mime=safe_upload(upload)
            except Exception as exc: st.error(str(exc)); return
            execute("INSERT INTO documents(title,category,description,file_name,file_data,file_mime,created_by) VALUES(%s,%s,%s,%s,%s,%s,%s)",(title.strip(),category,description,name,data,mime,user["id"]))
            audit(user["id"],"upload_document","document",None,title.strip()); rerun()
    active,archive,deleted=st.tabs(["Активни","Архив","Изтрити"])
    for tab,where in ((active,"archived=FALSE AND is_deleted=FALSE"),(archive,"archived=TRUE AND is_deleted=FALSE"),(deleted,"is_deleted=TRUE")):
        with tab:
            for row in fetch_all(f"SELECT * FROM documents WHERE {where} ORDER BY created_at DESC"):
                c1,c2,c3=st.columns([5,1,1]); c1.write(f"**{row['title']}** · {row['category']}")
                if not row["is_deleted"]:
                    if c2.button("Архив/Върни",key=f"document_archive_{row['id']}"):
                        execute("UPDATE documents SET archived=NOT archived,archived_at=CASE WHEN archived THEN NULL ELSE NOW() END WHERE id=%s",(row["id"],)); rerun()
                    if c3.button("Изтрий",key=f"document_delete_{row['id']}"):
                        execute("UPDATE documents SET is_deleted=TRUE,archived=TRUE,deleted_at=NOW() WHERE id=%s",(row["id"],)); audit(user["id"],"soft_delete_document","document",row["id"]); rerun()
                elif c2.button("Възстанови",key=f"document_restore_{row['id']}"):
                    execute("UPDATE documents SET is_deleted=FALSE,deleted_at=NULL WHERE id=%s",(row["id"],)); rerun()


def admin_polls_tab(user):
    st.subheader("Ново гласуване")
    title=st.text_input("Заглавие",key="poll_title")
    question=st.text_input("Въпрос",key="poll_question")
    description=st.text_area("Описание",key="poll_description")
    mode=st.radio("Варианти",["Да / Не","Персонализирани опции"],horizontal=True)
    custom=st.text_area("Персонализирани опции, по една на ред",disabled=mode=="Да / Не",key="poll_custom_options")
    c1,c2=st.columns(2)
    start_date=c1.date_input("Начална дата",date.today(),key="poll_start_date")
    start_time=c1.time_input("Начален час",time(0,0),key="poll_start_time")
    end_date=c2.date_input("Крайна дата",date.today(),key="poll_end_date")
    end_time=c2.time_input("Краен час",time(23,59),key="poll_end_time")
    files=st.file_uploader("Файлове и изображения",accept_multiple_files=True,key="poll_files")
    if st.button("Създай гласуване",type="primary",use_container_width=True):
        options=["Да","Не"] if mode=="Да / Не" else [x.strip() for x in custom.splitlines() if x.strip()]
        starts=datetime.combine(start_date,start_time); ends=datetime.combine(end_date,end_time)
        if not title.strip() or not question.strip() or len(options)<2 or ends<=starts:
            st.error("Попълнете заглавие, въпрос, поне две опции и валиден период.")
        else:
            try:
                attachments=[safe_upload(f) for f in files]
            except Exception as exc: st.error(str(exc)); return
            poll=execute("INSERT INTO polls(title,question,description,starts_at,ends_at,created_by) VALUES(%s,%s,%s,%s,%s,%s) RETURNING id",(title.strip(),question.strip(),description,starts,ends,user["id"]),True)
            for order,option in enumerate(options): execute("INSERT INTO poll_options(poll_id,option_text,sort_order) VALUES(%s,%s,%s)",(poll["id"],option,order))
            for name,data,mime in attachments:
                execute("INSERT INTO poll_attachments(poll_id,file_name,file_data,file_mime,is_image) VALUES(%s,%s,%s,%s,%s)",(poll["id"],name,data,mime,bool(mime and mime.startswith("image/"))))
            audit(user["id"],"create_poll","poll",poll["id"],title.strip()); rerun()
    st.subheader("Гласувания и резултати")
    polls=fetch_all("SELECT * FROM polls ORDER BY created_at DESC")
    if not polls: st.info("Няма създадени гласувания.")
    for poll in polls:
        now=datetime.now(poll["ends_at"].tzinfo) if poll["ends_at"].tzinfo else datetime.now()
        state="Активно" if poll["active"] and poll["starts_at"]<=now<=poll["ends_at"] else "Приключило / архив"
        with st.expander(f"{state} · {poll['title']}"):
            total=fetch_one("SELECT COUNT(*) AS count FROM poll_votes WHERE poll_id=%s",(poll["id"],))["count"]
            st.write(f"**{poll['question']}**"); st.caption(f"Гласували апартаменти: {total}")
            results=fetch_all("""SELECT po.option_text,COUNT(pv.id) AS votes FROM poll_options po
                LEFT JOIN poll_votes pv ON pv.option_id=po.id WHERE po.poll_id=%s
                GROUP BY po.id ORDER BY po.sort_order,po.id""",(poll["id"],))
            for result in results:
                percent=(result["votes"]/total*100) if total else 0
                st.write(f"**{result['option_text']}**: {result['votes']} гласа · {percent:.1f}%")
                st.progress(percent/100 if percent else 0)
            vote_details = fetch_all("""
                SELECT
                    a.apartment_number AS "Апартамент",
                    a.property_type AS "Тип обект",
                    CASE
                        WHEN pv.id IS NULL THEN 'Не е гласувал'
                        ELSE 'Гласувал'
                    END AS "Статус",
                    COALESCE(po.option_text, '') AS "Избран отговор",
                    pv.created_at AS "Дата и час",
                    CASE
                        WHEN pv.video_data IS NOT NULL THEN 'Видео'
                        WHEN pv.paper_signed = TRUE THEN 'Хартиен подпис'
                        WHEN pv.id IS NOT NULL THEN 'Очаква подпис'
                        ELSE ''
                    END AS "Подпис",
                    pv.id AS "_vote_id",
                    (pv.video_data IS NOT NULL) AS "_has_video",
                    pv.paper_signed AS "_paper_signed"
                FROM apartments a
                LEFT JOIN poll_votes pv
                    ON pv.apartment_id = a.id
                   AND pv.poll_id = %s
                LEFT JOIN poll_options po
                    ON po.id = pv.option_id
                WHERE a.active = TRUE
                ORDER BY
                    CASE
                        WHEN a.apartment_number ~ '^[0-9]+$'
                        THEN a.apartment_number::INTEGER
                        ELSE 999999
                    END,
                    a.apartment_number
            """, (poll["id"],))

            st.markdown("#### Доказателства за гласуване")
            voted_rows = [dict(row) for row in vote_details if row.get("_vote_id")]
            if not voted_rows:
                st.info("Все още няма подадени гласове.")
            for vote_row in voted_rows:
                with st.expander(
                    f"Ап. {vote_row['Апартамент']} · {vote_row['Избран отговор']} · {vote_row['Подпис']}"
                ):
                    if vote_row.get("_has_video"):
                        video_row=fetch_one("SELECT video_data,video_mime,video_name FROM poll_votes WHERE id=%s",(vote_row["_vote_id"],))
                        if video_row and video_row.get("video_data"):
                            st.video(bytes(video_row["video_data"]),format=video_row.get("video_mime") or "video/mp4")
                    elif vote_row.get("_paper_signed"):
                        st.success("Хартиеният подпис е отбелязан като получен.")
                    else:
                        st.warning("Очаква се хартиен подпис.")
                        if st.button("✍️ Отбележи получен подпис",key=f"paper_sign_{poll['id']}_{vote_row['_vote_id']}"):
                            execute("""UPDATE poll_votes SET proof_type='paper',paper_signed=TRUE,
                                paper_signed_at=NOW(),paper_signed_by=%s WHERE id=%s""",(user["id"],vote_row["_vote_id"]))
                            audit(user["id"],"confirm_paper_vote_signature","poll_vote",vote_row["_vote_id"],str(poll["id"]))
                            rerun()

            vote_df = pd.DataFrame(
                [dict(row) for row in vote_details]
            )

            if not vote_df.empty:
                vote_df = vote_df.drop(
                    columns=["_vote_id", "_has_video", "_paper_signed"],
                    errors="ignore",
                )
            if not vote_df.empty and "Дата и час" in vote_df.columns:
                vote_df["Дата и час"] = pd.to_datetime(
                    vote_df["Дата и час"],
                    errors="coerce",
                    utc=True
                ).dt.tz_localize(None)

            summary_df = pd.DataFrame([
                {
                    "Заглавие": poll["title"],
                    "Въпрос": poll["question"],
                    "Начало": poll["starts_at"],
                    "Край": poll["ends_at"],
                    "Общо гласували": int(total),
                    "Статус на анкетата": state,
                }
            ])

            for column in ["Начало", "Край"]:
                summary_df[column] = pd.to_datetime(
                    summary_df[column],
                    errors="coerce",
                    utc=True
                ).dt.tz_localize(None)

            results_df = pd.DataFrame([
                {
                    "Отговор": result["option_text"],
                    "Брой гласове": int(result["votes"]),
                    "Процент": round(
                        (result["votes"] / total * 100)
                        if total else 0,
                        2
                    ),
                }
                for result in results
            ])

            poll_export = io.BytesIO()

            with pd.ExcelWriter(
                poll_export,
                engine="openpyxl"
            ) as writer:

                summary_df.to_excel(
                    writer,
                    sheet_name="Информация",
                    index=False
                )

                results_df.to_excel(
                    writer,
                    sheet_name="Резултати",
                    index=False
                )

                vote_df.to_excel(
                    writer,
                    sheet_name="Апартаменти",
                    index=False
                )

                for sheet_name in writer.sheets:
                    worksheet = writer.sheets[sheet_name]
                    worksheet.freeze_panes = "A2"

                    for column_cells in worksheet.columns:
                        max_length = 0
                        column_letter = column_cells[0].column_letter

                        for cell in column_cells:
                            if cell.value is not None:
                                max_length = max(
                                    max_length,
                                    len(str(cell.value))
                                )

                        worksheet.column_dimensions[
                            column_letter
                        ].width = min(max_length + 3, 45)

            safe_poll_title = "".join(
                char if char.isalnum() or char in ("-", "_")
                else "_"
                for char in poll["title"]
            )

            st.download_button(
                "📥 Изтегли резултатите от анкетата",
                data=poll_export.getvalue(),
                file_name=f"poll_{poll['id']}_{safe_poll_title}.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                key=f"poll_export_{poll['id']}",
                use_container_width=True,
            )
            if st.button("Активирай/Деактивирай",key=f"poll_toggle_{poll['id']}"):
                execute("UPDATE polls SET active=NOT active WHERE id=%s",(poll["id"],)); rerun()

def _object_name(value):
    if isinstance(value,float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def _is_object_name(name):
    low=name.lower(); return name.isdigit() and 1<=int(name)<=59 or "офис" in low or low.startswith("пм")


def _number(value):
    if value is None or str(value).strip()=="": return None
    if isinstance(value,(int,float,Decimal)): return Decimal(str(value))
    import re
    m=re.search(r"-?\d+(?:[.,]\d+)?",str(value).replace(" ",""))
    return Decimal(m.group(0).replace(",",".")) if m else None


def apply_historical_snapshot(data,user_id):
    wb=load_workbook(io.BytesIO(data),data_only=True); ws=wb.worksheets[0]
    marker=next((r for r in range(1,ws.max_row+1) if str(ws.cell(r,3).value).strip()=="2026"),None)
    if not marker: raise ValueError("Не е намерен блокът 2026 в първия лист.")
    # The Excel snapshot is authoritative for 2024-2026. Clear old allocations for these years,
    # including test entries, then rebuild statuses exactly from the source.
    adjusted=fetch_one("SELECT COUNT(*) count FROM obligations WHERE year BETWEEN 2024 AND 2026 AND custom_status='force_majeure'")
    if adjusted and adjusted["count"]>0:
        raise ValueError("Има ръчни форсмажорни корекции за 2024–2026. Импортът е блокиран, за да не бъдат изтрити.")
    linked=fetch_all("""SELECT DISTINCT pa.payment_id FROM payment_allocations pa
        JOIN obligations o ON o.id=pa.obligation_id WHERE o.year BETWEEN 2024 AND 2026""")
    execute("DELETE FROM payment_allocations WHERE obligation_id IN (SELECT id FROM obligations WHERE year BETWEEN 2024 AND 2026)")
    for x in linked:
        execute("DELETE FROM payments WHERE id=%s AND NOT EXISTS (SELECT 1 FROM payment_allocations WHERE payment_id=%s)",(x["payment_id"],x["payment_id"]))
    execute("DELETE FROM obligations WHERE year BETWEEN 2024 AND 2026")
    stats={"objects":0,"paid_months":0,"fro_payments":0}
    for year,start_row,end_row in ((2025,2,marker-1),(2026,marker+1,ws.max_row)):
        currency="BGN" if year<2026 else "EUR"
        for row in range(start_row,end_row+1):
            raw=ws.cell(row,3).value
            if raw is None: continue
            name=_object_name(raw)
            if not _is_object_name(name): continue
            apt=fetch_one("SELECT id FROM apartments WHERE lower(apartment_number)=lower(%s)",(name,))
            if not apt:
                typ="Офис" if "офис" in name.lower() else "Паркомясто" if name.lower().startswith("пм") else "Апартамент"
                apt=execute("INSERT INTO apartments(apartment_number,property_type,adult_residents,annual_fro) VALUES(%s,%s,1,0) RETURNING id",(name,typ),True)
            apt_id=apt["id"]; stats["objects"]+=1
            vals=[ws.cell(row,c).value for c in range(4,16)]
            numeric=[_number(v) for v in vals if _number(v) is not None and _number(v)>0]
            default_fee=max(set(numeric),key=numeric.count) if numeric else Decimal("12" if year==2026 else "20")
            for month,cell in enumerate(vals,1):
                n=_number(cell); amount=n if n is not None and n>0 else default_fee
                ob=execute("""INSERT INTO obligations(apartment_id,fee_type,year,month,amount,note,source_currency)
                    VALUES(%s,'entrance',%s,%s,%s,'Исторически статус от Excel',%s)
                    ON CONFLICT(apartment_id,fee_type,year,month) DO UPDATE SET amount=EXCLUDED.amount,note=EXCLUDED.note,source_currency=EXCLUDED.source_currency RETURNING id""",(apt_id,year,month,amount,currency),True)
                filled=cell is not None and str(cell).strip()!=""
                if filled:
                    num=f"HIST-ENT-{year}-{apt_id}-{month}"
                    pay=execute("""INSERT INTO payments(payment_number,apartment_id,amount,payment_method,status,payment_date,note,created_by,confirmed_by,confirmed_at)
                        VALUES(%s,%s,%s,'cash','confirmed',%s,'Исторически платено по Excel',%s,%s,NOW()) RETURNING id""",(num,apt_id,amount,date(year,month,1),user_id,user_id),True)
                    execute("UPDATE payments SET source_currency=%s WHERE id=%s",(currency,pay["id"]))
                    if amount>0: execute("INSERT INTO payment_allocations(payment_id,obligation_id,amount) VALUES(%s,%s,%s)",(pay["id"],ob["id"],amount))
                    stats["paid_months"]+=1
            annual=_number(ws.cell(row,18).value) or Decimal("0"); fro_paid=_number(ws.cell(row,16).value) or Decimal("0")
            if annual>0:
                monthly=(annual/Decimal("12")).quantize(Decimal("0.01")); obs=[]
                for month in range(1,13):
                    amount=annual-monthly*11 if month==12 else monthly
                    ob=execute("""INSERT INTO obligations(apartment_id,fee_type,year,month,amount,note,source_currency)
                        VALUES(%s,'fro',%s,%s,%s,'Исторически ФРО от Excel',%s)
                        ON CONFLICT(apartment_id,fee_type,year,month) DO UPDATE SET amount=EXCLUDED.amount,note=EXCLUDED.note,source_currency=EXCLUDED.source_currency RETURNING id""",(apt_id,year,month,amount,currency),True); obs.append((ob["id"],amount))
                remaining=min(fro_paid,annual)
                if remaining>0:
                    pay=execute("""INSERT INTO payments(payment_number,apartment_id,amount,payment_method,status,payment_date,note,created_by,confirmed_by,confirmed_at)
                        VALUES(%s,%s,%s,'cash','confirmed',%s,'Исторически ФРО по Excel',%s,%s,NOW()) RETURNING id""",(f"HIST-FRO-{year}-{apt_id}",apt_id,remaining,date(year,1,1),user_id,user_id),True)
                    execute("UPDATE payments SET source_currency=%s WHERE id=%s",(currency,pay["id"]))
                    for oid,amount in obs:
                        part=min(amount,remaining)
                        if part>0: execute("INSERT INTO payment_allocations(payment_id,obligation_id,amount) VALUES(%s,%s,%s)",(pay["id"],oid,part)); remaining-=part
                    stats["fro_payments"]+=1
    # 2024: only these remain unpaid; all other 2024 legacy balances are treated as settled.
    entrance={"36":120,"42":240,"45":240,"46":1,"47":1,"48":240,"19":240,"52":120,"29":240,"58":240,"59":240}
    fro={"15":121.80,"20":129.18,"29":127.95,"30":115.65,"35":136.57,"36":125.49,"42":134.00,"49":134.00,"52":170.00,"53":130.00,"58":193.00,"59":215.00}
    execute("DELETE FROM obligations WHERE year=2024 AND note IN ('Старо задължение от Excel','Остатък ФРО 2024')")
    for name,amount in entrance.items():
        apt=fetch_one("SELECT id FROM apartments WHERE lower(apartment_number)=lower(%s)",(name,))
        if apt: execute("INSERT INTO obligations(apartment_id,fee_type,year,month,amount,note,source_currency) VALUES(%s,'old_debt',2024,0,%s,'Старо задължение от Excel','BGN') ON CONFLICT(apartment_id,fee_type,year,month) DO UPDATE SET amount=EXCLUDED.amount,note=EXCLUDED.note,source_currency='BGN'",(apt["id"],amount))
    for name,amount in fro.items():
        apt=fetch_one("SELECT id FROM apartments WHERE lower(apartment_number)=lower(%s)",(name,))
        if apt:
            execute("DELETE FROM obligations WHERE apartment_id=%s AND fee_type='fro' AND year=2024",(apt["id"],))
            execute("INSERT INTO obligations(apartment_id,fee_type,year,month,amount,note,source_currency) VALUES(%s,'fro',2024,0,%s,'Остатък ФРО 2024','BGN')",(apt["id"],amount))
    # Baseline is written after history, so historical rows do not change current cash balances.
    execute("DELETE FROM balance_baseline")
    execute("INSERT INTO balance_baseline(cash_amount,bank_amount,applied_at,applied_by,note) VALUES(2398.13,2096.53,NOW(),%s,'Начални наличности от Excel')",(user_id,))
    audit(user_id,"apply_historical_snapshot","system",None,str(stats)); return stats

def admin_panel(user):
    st.title("Административен панел")
    admin_sections = [
        "Табло", "Апартаменти", "Регистрации", "Роли",
        "Плащания", "Исторически статус", "Експорт",
        "Технически проблеми", "Обяви", "Документи",
        "Гласувания", "Паркоместа",
    ]
    admin_section = st.selectbox(
        "Административен модул",
        admin_sections,
        key="admin_active_section",
    )
    st.divider()
    if admin_section == "Табло":
        dashboard_base = fetch_one("""
            SELECT *
            FROM dashboard_baseline
            ORDER BY applied_at DESC
            LIMIT 1
        """)

        if not dashboard_base:
            dashboard_base = {
                "expense_amount": Decimal("7000.00"),
                "applied_at": datetime.now(),
            }

        s = fetch_one("""
            SELECT
                (SELECT COUNT(*)
                 FROM apartments
                 WHERE active) AS a,

                (SELECT COUNT(*)
                 FROM users
                 WHERE account_status = 'pending') AS p,

                (SELECT COALESCE(SUM(amount), 0)
                 FROM payments
                 WHERE status = 'confirmed'
                   AND source_currency = 'EUR') AS income,

                (SELECT COALESCE(SUM(amount), 0)
                 FROM expenses
                 WHERE created_at > %s) AS new_expenses
        """, (dashboard_base["applied_at"],))

        total_income = Decimal(s["income"] or 0)

        total_expenses = (
            Decimal(dashboard_base["expense_amount"] or 0)
            + Decimal(s["new_expenses"] or 0)
        )

        c1, c2, c3, c4 = st.columns(4)

        c1.metric(
            "Обекти",
            s["a"]
        )

        c2.metric(
            "Чакащи",
            s["p"]
        )

        c3.metric(
            "Постъпления",
            fmt(total_income)
        )

        c4.metric(
            "Разходи",
            fmt(total_expenses)
        )

        st.caption(
            "Разходите започват от начална стойност "
            f"{fmt(dashboard_base['expense_amount'])} и към нея "
            "се добавят новите разходи след зануляването."
        )
    if admin_section == "Апартаменти":
        st.subheader("Добавяне на нов обект")
        with st.form("apt"):
            c1,c2,c3=st.columns(3); num=c1.text_input("Номер/обект"); typ=c2.selectbox("Тип",["Апартамент","Офис","Магазин","Гараж","Паркомясто"]); floor=c3.text_input("Етаж")
            adults=st.number_input("Лица над 18 г.",0,20,1); fro=st.number_input("Годишен ФРО",0.0,10000.0,0.0,.01); add=st.form_submit_button("Добави")
        if add and num:
            try: execute("INSERT INTO apartments(apartment_number,property_type,floor,adult_residents,annual_fro) VALUES(%s,%s,%s,%s,%s)",(num,typ,floor,adults,fro)); rerun()
            except Exception: st.error("Обектът вече съществува.")
        st.divider()
        st.subheader("Корекция на хора, такса вход и годишен ФРО")

        all_apts = fetch_all("""
            SELECT
                id,
                apartment_number,
                property_type,
                adult_residents,
                monthly_entrance_fee,
                annual_fro,
                admin_note
            FROM apartments
            WHERE active
            ORDER BY apartment_number
        """)

        edit_map = {
            f"{x['apartment_number']} · {x['property_type']}": x
            for x in all_apts
        }

        if edit_map:
            selected_key = st.selectbox(
                "Избери обект за корекция",
                list(edit_map),
                key="edit_apartment"
            )
            current = edit_map[selected_key]
            current_adults = int(current["adult_residents"] or 0)
            current_entrance_fee = (
                Decimal(current["monthly_entrance_fee"])
                if current.get("monthly_entrance_fee") is not None
                else Decimal(current_adults) * Decimal("6.00")
            )
            current_fro = Decimal(current["annual_fro"] or 0)

            st.caption(
                f"Текущо: {current_adults} лица · "
                f"Такса вход {current_entrance_fee:.2f} € / месец · "
                f"ФРО {current_fro:.2f} € / година"
            )

            with st.form(f"edit_apartment_form_{current['id']}"):
                ec1, ec2, ec3 = st.columns(3)
                new_adults = ec1.number_input(
                    "Лица над 18 г.",
                    min_value=0,
                    max_value=20,
                    value=current_adults,
                    step=1,
                    key=f"edit_adults_{current['id']}"
                )
                new_entrance_fee = ec2.number_input(
                    "Месечна такса вход (€)",
                    min_value=0.0,
                    max_value=10000.0,
                    value=float(current_entrance_fee),
                    step=0.01,
                    format="%.2f",
                    key=f"edit_entrance_fee_{current['id']}"
                )
                new_fro = ec3.number_input(
                    "Годишен ФРО (€)",
                    min_value=0.0,
                    max_value=10000.0,
                    value=float(current_fro),
                    step=0.01,
                    format="%.2f",
                    key=f"edit_fro_{current['id']}"
                )
                admin_note = st.text_input(
                    "Коментар / бележка",
                    value=str(current.get("admin_note") or ""),
                    placeholder="Напр. Под наем / Собственик на ап. 2 и 3 / Необитаем",
                    key=f"edit_admin_note_{current['id']}"
                )
                effective_year = st.selectbox(
                    "Преизчисли неплатените задължения за година",
                    YEARS,
                    index=YEARS.index(date.today().year) if date.today().year in YEARS else 0,
                    key=f"edit_year_{current['id']}"
                )
                st.info(
                    "Ще бъдат променени само напълно неплатените такси вход и ФРО "
                    "за избраната година. Платените и частично платените позиции се запазват."
                )
                save_edit = st.form_submit_button(
                    "Запази корекцията",
                    use_container_width=True
                )

            if save_edit:
                entrance_fee = Decimal(str(new_entrance_fee)).quantize(Decimal("0.01"))
                annual_fro = Decimal(str(new_fro)).quantize(Decimal("0.01"))

                execute("""
                    UPDATE apartments
                    SET adult_residents = %s,
                        monthly_entrance_fee = %s,
                        annual_fro = %s,
                        admin_note = %s
                    WHERE id = %s
                """, (
                    new_adults,
                    entrance_fee,
                    annual_fro,
                    admin_note.strip(),
                    current["id"]
                ))

                execute("""
                    INSERT INTO resident_history(
                        apartment_id, adults, effective_from, reason, created_by
                    )
                    VALUES(%s, %s, %s, 'Административна корекция', %s)
                    ON CONFLICT(apartment_id, effective_from)
                    DO UPDATE SET
                        adults = EXCLUDED.adults,
                        reason = EXCLUDED.reason,
                        created_by = EXCLUDED.created_by
                """, (
                    current["id"],
                    new_adults,
                    date(int(effective_year), 1, 1),
                    user["id"]
                ))

                execute("""
                    UPDATE obligations o
                    SET amount = %s
                    WHERE o.apartment_id = %s
                      AND o.year = %s
                      AND o.fee_type = 'entrance'
                      AND o.custom_status IS DISTINCT FROM 'force_majeure'
                      AND COALESCE(
                            (
                                SELECT SUM(pa.amount)
                                FROM payment_allocations pa
                                JOIN payments p ON p.id = pa.payment_id
                                WHERE pa.obligation_id = o.id
                                  AND p.status = 'confirmed'
                            ),
                            0
                          ) = 0
                """, (entrance_fee, current["id"], effective_year))

                monthly_fro = (annual_fro / Decimal("12")).quantize(Decimal("0.01"))
                for month in range(1, 13):
                    fro_amount = annual_fro - monthly_fro * 11 if month == 12 else monthly_fro
                    execute("""
                        UPDATE obligations o
                        SET amount = %s
                        WHERE o.apartment_id = %s
                          AND o.year = %s
                          AND o.month = %s
                          AND o.fee_type = 'fro'
                          AND COALESCE(
                                (
                                    SELECT SUM(pa.amount)
                                    FROM payment_allocations pa
                                    JOIN payments p ON p.id = pa.payment_id
                                    WHERE pa.obligation_id = o.id
                                      AND p.status = 'confirmed'
                                ),
                                0
                              ) = 0
                    """, (fro_amount, current["id"], effective_year, month))

                audit(
                    user["id"],
                    "update_apartment_financial_settings",
                    "apartment",
                    current["id"],
                    f"adults={new_adults};entrance_fee={entrance_fee};fro={annual_fro};year={effective_year}"
                )
                st.success(
                    "Корекцията е записана. Обновени са хората, месечната такса вход "
                    "и годишният ФРО. Платените позиции са запазени."
                )
                rerun()

        st.dataframe(
            fetch_all("""
                SELECT
                    apartment_number AS обект,
                    property_type AS тип,
                    adult_residents AS лица,
                    COALESCE(
                        monthly_entrance_fee,
                        adult_residents * 6
                    ) AS "Месечна такса",
                    annual_fro AS "Годишен ФРО",
                    status AS статус,
                    COALESCE(admin_note, '') AS коментар
                FROM apartments
                ORDER BY apartment_number
            """),
            use_container_width=True,
            hide_index=True
        )
    if admin_section == "Регистрации":
        regs=fetch_all("SELECT u.id,u.first_name,u.last_name,u.email,a.apartment_number FROM users u LEFT JOIN apartments a ON a.id=u.apartment_id WHERE u.account_status='pending'")
        for x in regs:
            c1,c2,c3=st.columns([4,1,1]); c1.write(f"**{x['first_name']} {x['last_name']}** · ап. {x['apartment_number']} · {x['email']}")
            if c2.button("Одобри",key=f"ok{x['id']}"): execute("UPDATE users SET account_status='approved' WHERE id=%s",(x["id"],)); rerun()
            if c3.button("Откажи",key=f"no{x['id']}"): execute("UPDATE users SET account_status='rejected' WHERE id=%s",(x["id"],)); rerun()
    if admin_section == "Роли":
        users=fetch_all("SELECT id,first_name,last_name,email,role FROM users ORDER BY first_name,last_name")
        for x in users:
            c1,c2,c3=st.columns([4,2,1]); c1.write(f"**{x['first_name']} {x['last_name']}** · {x['email']}"); roles=["neighbor","cashier","manager"]; role=c2.selectbox("Роля",roles,index=roles.index(x["role"]) if x["role"] in roles else 0,key=f"role{x['id']}",label_visibility="collapsed")
            if c3.button("Запази",key=f"rs{x['id']}"): execute("UPDATE users SET role=%s,account_status='approved' WHERE id=%s",(role,x["id"])); audit(user["id"],"change_role","user",x["id"],role); rerun()
    if admin_section == "Плащания":

        st.subheader("Плащания")
    
        rows = fetch_all("""
            SELECT
                p.id,
                p.created_at AS дата,
                a.apartment_number AS обект,
                p.payment_number AS номер,
                p.amount AS сума,
                p.payment_method AS метод,
                p.status AS статус,
                p.provider_transaction_id AS референция
            FROM payments p
            JOIN apartments a ON a.id = p.apartment_id
            ORDER BY p.created_at DESC
            LIMIT 150
        """)
    
        if not rows:
            st.info("Няма плащания.")
        else:
    
            for row in rows:
    
                c1, c2 = st.columns([10, 1])
    
                with c1:
    
                    st.markdown(
                        f"""
    **Обект:** {row['обект']} |
    **Номер:** {row['номер']} |
    **Сума:** {row['сума']} |
    **Метод:** {method_bg(row['метод'])} |
    **Статус:** {payment_status_bg(row['статус'])}
    """
                    )
    
                with c2:
    
                    if row["статус"] == "confirmed":
    
                        if st.button(
                            "↩️ Отмени",
                            key=f"undo_payment_{row['id']}"
                        ):
    
                            execute(
                                """
                                DELETE FROM payment_allocations
                                WHERE payment_id=%s
                                """,
                                (row["id"],)
                            )
    
                            execute(
                                """
                                UPDATE payments
                                SET status='cancelled'
                                WHERE id=%s
                                """,
                                (row["id"],)
                            )
    
                            audit(
                                user["id"],
                                "cancel_payment",
                                "payment",
                                row["id"],
                                row["номер"]
                            )
    
                            st.success(
                                f"Плащането {row['номер']} е отменено."
                            )
                            st.warning(
                                "Задължението е върнато автоматично към длъжника."
                            )
    
                            rerun()
    
                st.divider()
    
        st.subheader("История на форсмажорни корекции")
    
        adjustments = fetch_all("""
            SELECT
                oa.changed_at AS дата,
                a.apartment_number AS обект,
                o.year AS година,
                o.month AS месец,
                o.fee_type AS вид,
                oa.old_amount AS стара_сума,
                oa.new_amount AS нова_сума,
                oa.action AS действие,
                oa.reason AS основание
            FROM obligation_adjustments oa
            JOIN obligations o ON o.id = oa.obligation_id
            JOIN apartments a ON a.id = o.apartment_id
            ORDER BY oa.changed_at DESC
            LIMIT 200
        """)
    
        st.dataframe(
            adjustments,
            use_container_width=True,
            hide_index=True
        )
    
    if admin_section == "Исторически статус":
        st.subheader("Исторически статус 2024–2026")
        st.info("Качи оригиналния Excel. Първият лист задава платените месеци и ФРО за 2025 и 2026. Задълженията за 2024 и началните каси са зададени по предоставения списък.")
        file=st.file_uploader("Excel файл",type=["xlsx"],key="history_excel")
        confirm=st.checkbox("Потвърждавам: начална каса 2 398,13 €, банка 2 096,53 €")
        if st.button("Приложи реалния исторически статус",type="primary",disabled=not(file and confirm)):
            try:
                stats=apply_historical_snapshot(file.getvalue(),user["id"]); st.success(f"Готово: {stats['paid_months']} платени месеца и {stats['fro_payments']} ФРО плащания."); rerun()
            except Exception as exc: st.error(f"Миграцията не беше приложена: {exc}")
        base=fetch_one("SELECT * FROM balance_baseline ORDER BY applied_at DESC LIMIT 1")
        if base: st.success(f"Активно начално салдо: в брой {fmt(base['cash_amount'])}, банка {fmt(base['bank_amount'])}")
    if admin_section == "Експорт":
        rows=fetch_all("SELECT a.apartment_number,o.fee_type,o.year,o.month,o.amount,COALESCE(SUM(CASE WHEN p.status='confirmed' THEN pa.amount ELSE 0 END),0) paid FROM obligations o JOIN apartments a ON a.id=o.apartment_id LEFT JOIN payment_allocations pa ON pa.obligation_id=o.id LEFT JOIN payments p ON p.id=pa.payment_id GROUP BY a.apartment_number,o.id ORDER BY a.apartment_number,o.year,o.month")
        out=io.BytesIO(); pd.DataFrame(rows).to_excel(out,index=False,engine="openpyxl"); st.download_button("Изтегли общ отчет",out.getvalue(),"Building_Manager_Report.xlsx")
    if admin_section == "Технически проблеми": admin_technical_tab(user)
    if admin_section == "Обяви": admin_announcements_tab(user)
    if admin_section == "Документи": admin_documents_tab(user)
    if admin_section == "Гласувания": admin_polls_tab(user)
    if admin_section == "Паркоместа": parking_page()

user=current_user()
if not user:
    login_page()
    show_email_warning()
    st.stop()

show_email_warning()

if user["role"] == "neighbor":
    if "neighbor_page" not in st.session_state:
        st.session_state.neighbor_page = "Моите задължения"
    with st.sidebar:
        if st.button("Изход", use_container_width=True, key="neighbor_logout"):
            logout_user()
            st.session_state.pop("neighbor_page", None)
            rerun()
    top_col, spacer_col = st.columns([1, 9])
    with top_col:
        neighbor_notification_bell(user)
    neighbor_main_navigation()
    page = st.session_state.neighbor_page
else:
    with st.sidebar:
        st.title("🏢 Building.Manager")
        st.caption(f"{user['first_name']} {user['last_name']} | {user['role']}")
        pages = ["Каса", "Отчет", "Разходи", "Админ панел"] if user["role"] == "manager" else ["Каса", "Отчет", "Разходи"]
        page = st.radio("Навигация", pages)
        if st.button("Изход", use_container_width=True, key="staff_logout"):
            logout_user()
            rerun()

if page=="Каса": cashier_page(user)
elif page=="Отчет": report_page(user)
elif page=="Разходи": expenses_page(user)
elif page=="Моите задължения": neighbor_page(user)
elif page=="Архив": neighbor_archive(user)
elif page=="Технически проблем": technical_issues_page(user)
elif page=="Обяви": announcements_page()
elif page=="Документи": documents_page()
elif page=="Електронно гласуване": voting_page(user)
elif page=="Паркоместа": parking_page()
elif page=="Админ панел": admin_panel(user)
